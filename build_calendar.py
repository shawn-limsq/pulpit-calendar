#!/usr/bin/env python3
"""Build the pulpit-schedule calendar feed from the newest file in schedule/.

Supports .docx, .pdf, .xlsx and .csv. Columns are matched by their header
names, so a new format with extra columns (e.g. Chairperson) works as long as
the header row has recognisable names. Every event is an all-day event.

Usage: python build_calendar.py [path/to/schedule-file]
Output: docs/preaching.ics
"""
import csv
import hashlib
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from icalendar import Calendar, Event

ROOT = Path(__file__).resolve().parent
SCHEDULE_DIR = ROOT / "schedule"
OUT = ROOT / "docs" / "preaching.ics"
CAL_NAME = "LBPC Pulpit Schedule"
UID_DOMAIN = "pulpit.lifebpc.com"

# Header names (lower-case, substring match) -> field
COLUMN_ALIASES = {
    "date": ["date"],
    "text": ["text", "passage", "scripture", "reading"],
    "title": ["title", "sermon", "topic", "theme"],
    "speaker": ["speaker", "preacher"],
    "chair": ["chair", "presid", "worship leader", "liturgist"],
    "notes": ["notes", "remarks", "description"],
    "movement": ["movement", "events"],
}

SCRIPTURE_RE = re.compile(r"\b\d+\s*:\s*\d+")
DATE_FORMATS = ["%d %b %y", "%d %b %Y", "%d %B %y", "%d %B %Y", "%d/%m/%Y",
                "%d/%m/%y", "%Y-%m-%d", "%d-%b-%y", "%d-%b-%Y", "%d.%m.%Y"]


def clean(s):
    s = (s or "").replace(" ", " ")
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in s.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def parse_date(s):
    s = re.sub(r"\(.*?\)", "", clean(s)).strip()
    s = re.sub(r"^(mon|tue|wed|thu|fri|sat|sun)\w*,?\s*", "", s, flags=re.I)
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


# ---------- readers: each returns a list of rows (list of cell strings) ----------

def rows_from_docx(path):
    import docx
    rows = []
    for table in docx.Document(path).tables:
        for r in table.rows:
            rows.append([c.text for c in r.cells])
    return rows


def rows_from_pdf(path):
    import pdfplumber
    rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend([[c or "" for c in r] for r in table])
    return rows


def rows_from_xlsx(path):
    import openpyxl
    rows = []
    for ws in openpyxl.load_workbook(path, data_only=True).worksheets:
        for r in ws.iter_rows(values_only=True):
            rows.append([v.strftime("%d %b %Y") if isinstance(v, (datetime, date))
                         else ("" if v is None else str(v)) for v in r])
    return rows


def rows_from_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


READERS = {".docx": rows_from_docx, ".pdf": rows_from_pdf,
           ".xlsx": rows_from_xlsx, ".csv": rows_from_csv}


# ---------- interpretation ----------

def map_header(row):
    cells = [clean(c).lower() for c in row]
    if not any(c == "date" or c.startswith("date") for c in cells):
        return None
    mapping = {}
    for i, c in enumerate(cells):
        for field, aliases in COLUMN_ALIASES.items():
            if field not in mapping and any(a in c for a in aliases):
                mapping[field] = i
                break
    return mapping if "date" in mapping else None


def extract_services(rows):
    services, colmap, series = [], None, ""
    for row in rows:
        header = map_header(row)
        if header:
            colmap = header
            continue
        if not colmap:
            continue
        get = lambda f: clean(row[colmap[f]]) if f in colmap and colmap[f] < len(row) else ""
        d = parse_date(get("date"))
        if not d:
            # A merged full-width row is a series / quarter heading
            vals = {clean(c) for c in row if clean(c)}
            if len(vals) == 1:
                series = vals.pop()
            continue
        svc = {f: get(f) for f in COLUMN_ALIASES if f != "date"}
        svc["date"], svc["series"] = d, series
        # Passage sometimes lands in Notes when Text is blank
        if not svc["text"] and SCRIPTURE_RE.search(svc["notes"]) and len(svc["notes"]) < 40:
            svc["text"], svc["notes"] = svc["notes"], ""
        services.append(svc)
    return services


def build_ics(services, source_name):
    cal = Calendar()
    cal.add("prodid", "-//LBPC//Pulpit Schedule//EN")
    cal.add("version", "2.0")
    cal.add("calscale", "GREGORIAN")
    cal.add("method", "PUBLISH")
    cal.add("x-wr-calname", CAL_NAME)
    cal.add("x-wr-caldesc", f"Pulpit schedule (source: {source_name}). Updates automatically.")
    cal.add("refresh-interval", timedelta(hours=1), parameters={"VALUE": "DURATION"})
    cal["x-published-ttl"] = "PT1H"

    stamp = datetime.now(timezone.utc).replace(microsecond=0)
    seen = {}
    for s in sorted(services, key=lambda s: s["date"]):
        n = seen.get(s["date"], 0)
        seen[s["date"]] = n + 1
        # UID depends only on the date (and order within the date), so edits
        # to speaker/title update the existing event instead of duplicating it.
        uid_key = s["date"].isoformat() + (f"-{n}" if n else "")
        uid = hashlib.sha1(uid_key.encode()).hexdigest()[:16] + "@" + UID_DOMAIN

        title_line = s["title"].replace("\n", " – ")
        summary = f"{s['speaker'] or 'TBA'}: {title_line}" if title_line else (s["speaker"] or "Service")
        if s["text"]:
            summary += f" ({s['text'].splitlines()[0]})"

        desc = []
        for label, key in [("Speaker", "speaker"), ("Chairperson", "chair"),
                           ("Passage", "text"), ("Sermon", "title"),
                           ("Series", "series"), ("Also", "movement")]:
            if s.get(key):
                desc.append(f"{label}: {s[key].replace(chr(10), ' – ')}")
        if s["notes"]:
            desc.append("\n" + s["notes"])

        ev = Event()
        ev.add("uid", uid)
        ev.add("dtstamp", stamp)
        ev.add("dtstart", s["date"])
        ev.add("dtend", s["date"] + timedelta(days=1))
        ev.add("summary", summary)
        ev.add("description", "\n".join(desc))
        ev.add("transp", "TRANSPARENT")  # all-day, doesn't block time
        cal.add_component(ev)
    return cal.to_ical()


def _changed_at(p):
    """Last commit time of the file (reliable on GitHub), else file mtime."""
    import subprocess
    try:
        out = subprocess.run(["git", "log", "-1", "--format=%ct", "--", str(p)],
                             cwd=ROOT, capture_output=True, text=True).stdout.strip()
        if out:
            return int(out)
    except OSError:
        pass
    return p.stat().st_mtime


def newest_schedule():
    files = [p for p in SCHEDULE_DIR.iterdir() if p.suffix.lower() in READERS]
    if not files:
        sys.exit("No schedule file found in schedule/")
    return max(files, key=_changed_at)


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else newest_schedule()
    rows = READERS[path.suffix.lower()](path)
    services = extract_services(rows)
    if not services:
        sys.exit(f"Found no dated rows in {path.name}; check the header row has a 'Date' column.")
    new = build_ics(services, path.name)
    strip = lambda b: re.sub(rb"DTSTAMP:[^\r\n]*", b"", b)
    if OUT.exists() and strip(OUT.read_bytes()) == strip(new):
        print(f"{len(services)} services from {path.name}; no changes to the calendar.")
        return
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_bytes(new)
    print(f"{len(services)} services from {path.name} -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
